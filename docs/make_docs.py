"""Generates the documentation as Word (.docx, requires python-docx) and Markdown (.md) files.

    python docs/make_docs.py

Both formats are written from the same content, so they always match.
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent


class WordDoc:
    def __init__(self, title, subtitle):
        self.doc = Document()
        style = self.doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        heading = self.doc.add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = heading.add_run(title)
        run.bold = True
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        sub = self.doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub.add_run(subtitle)
        sub_run.font.size = Pt(12)
        sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    def h1(self, text):
        self.doc.add_heading(text, level=1)

    def h2(self, text):
        self.doc.add_heading(text, level=2)

    def h3(self, text):
        self.doc.add_heading(text, level=3)

    def page_break(self):
        self.doc.add_page_break()

    def p(self, text, bold_prefix=None):
        para = self.doc.add_paragraph()
        if bold_prefix:
            para.add_run(bold_prefix).bold = True
        para.add_run(text)
        return para

    def bullets(self, items, style="List Bullet"):
        for item in items:
            para = self.doc.add_paragraph(style=style)
            if isinstance(item, tuple):
                para.add_run(item[0]).bold = True
                para.add_run(item[1])
            else:
                para.add_run(item)

    def steps(self, items):
        """Numbered list that restarts at 1 (a new numbering instance per list)."""
        numbering = self.doc.part.numbering_part.numbering_definitions._numbering
        base_id = self.doc.styles["List Number"]._element.pPr.numPr.numId.val
        abstract_id = numbering.num_having_numId(base_id).abstractNumId.val
        num = numbering.add_num(abstract_id)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        for item in items:
            para = self.doc.add_paragraph(style="List Number")
            num_pr = para._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = num.numId
            para.add_run(item)

    def code(self, text):
        for line in text.strip("\n").splitlines():
            para = self.doc.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.left_indent = Pt(18)
            run = para.add_run(line)
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def table(self, header, rows, widths=None):
        table = self.doc.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, header):
            cell.text = ""
            cell.paragraphs[0].add_run(text).bold = True
            shading = OxmlElement("w:shd")
            shading.set(qn("w:val"), "clear")
            shading.set(qn("w:color"), "auto")
            shading.set(qn("w:fill"), "D9E2F3")
            cell._element.get_or_add_tcPr().append(shading)
        for row in rows:
            cells = table.add_row().cells
            for cell, text in zip(cells, row):
                cell.text = str(text)
        if widths:
            table.autofit = False
            for column, width in zip(table.columns, widths):
                column.width = Pt(width)
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Pt(width)
        self.doc.add_paragraph()

    def save(self, name):
        path = HERE / f"{name}.docx"
        self.doc.save(path)
        print(f"Written: {path}")


def md_escape(text):
    """Escape characters that Markdown would otherwise treat as markup."""
    return re.sub(r"([<*|`])", r"\\\1", str(text))


class MarkdownDoc:
    """Same interface as Doc, writes GitHub flavoured Markdown."""

    def __init__(self, title, subtitle):
        self.lines = [f"# {md_escape(title)}", "", f"_{md_escape(subtitle)}_", ""]

    def h1(self, text):
        self.lines += [f"## {md_escape(text)}", ""]

    def h2(self, text):
        self.lines += [f"### {md_escape(text)}", ""]

    def h3(self, text):
        self.lines += [f"#### {md_escape(text)}", ""]

    def page_break(self):
        self.lines += ["---", ""]

    def p(self, text, bold_prefix=None):
        prefix = f"**{md_escape(bold_prefix)}**" if bold_prefix else ""
        self.lines += [prefix + md_escape(text), ""]

    @staticmethod
    def _item(item):
        if isinstance(item, tuple):
            return f"**{md_escape(item[0].strip())}** {md_escape(item[1])}"
        return md_escape(item)

    def bullets(self, items):
        self.lines += [f"- {self._item(item)}" for item in items] + [""]

    def steps(self, items):
        self.lines += [f"{i}. {self._item(item)}" for i, item in enumerate(items, 1)] + [""]

    def code(self, text):
        text = text.strip("\n")
        language = "json" if text.lstrip().startswith("{") else "powershell"
        self.lines += [f"```{language}", text, "```", ""]

    def table(self, header, rows, widths=None):
        def row(cells):
            return "| " + " | ".join(md_escape(c) for c in cells) + " |"

        self.lines += [row(header), "|" + "---|" * len(header)] + [row(r) for r in rows] + [""]

    def save(self, name):
        path = HERE / f"{name}.md"
        path.write_text("\n".join(self.lines).rstrip() + "\n", encoding="utf-8")
        print(f"Written: {path}")


def technical_en(Doc):
    d = Doc("Trademark Image Checker", "Technical documentation - design, installation and operation (version 1.2.0)")

    d.h1("1. Purpose")
    d.p("Trademark Image Checker is a Windows desktop program that examines a raster image against a "
        "configurable set of criteria (file size, pixel dimensions, resolution and image type), shows "
        "the result in a GUI, lets the user choose valid target values for the failing properties and "
        "converts the image accordingly. Its default criteria correspond to typical trademark office "
        "image upload requirements. The user interface is localized (Hungarian by default, English included) "
        "and further languages can be added by editing a JSON file.")

    d.h1("2. Requirements")
    d.h2("2.1 Running the compiled program")
    d.bullets([
        "Windows 10 or Windows 11, 64-bit (x64).",
        "No Python installation is needed: dist\\TrademarkImage.exe is a self-contained single file.",
        "Write access to the folder of the converted images and, optionally, to the folder of the criteria file.",
    ])
    d.h2("2.2 Running from source")
    d.table(["Component", "Version", "Purpose"], [
        ["Python", "3.10 or newer (developed and tested with 3.14.0)", "Runtime; must include tkinter (Tk 8.6)"],
        ["Pillow", ">= 11.0 (tested 13.0.0)", "Image reading, conversion and writing"],
        ["tkinterdnd2", ">= 0.4 (tested 0.6.4)", "Drag and drop of files into the Tk window (bundles tkdnd 2.9/2.10)"],
        ["pytest", ">= 8.0 (tested 9.1.1)", "Automated tests (development only)"],
        ["PyInstaller", ">= 6.16 (tested 6.22.3)", "Building the Windows executable (development only)"],
        ["python-docx", ">= 1.1 (tested 1.2.0)", "Regenerating this documentation (development only)"],
    ], widths=[90, 150, 230])
    d.p("Runtime dependencies are listed in requirements.txt, development dependencies in requirements-dev.txt.")

    d.h2("2.3 Installation from source")
    d.code("""
git clone <repository> Trademarkimage
cd Trademarkimage
python -m venv .venv
.venv\\Scripts\\python.exe -m pip install -r requirements-dev.txt
""")
    d.h2("2.4 Running")
    d.code("""
.venv\\Scripts\\python.exe -m trademarkimage [--config FILE] [--lang CODE] [--locals FILE] [IMAGE]
.venv\\Scripts\\python.exe run_trademarkimage.py [options] [IMAGE]
dist\\TrademarkImage.exe [options] [IMAGE]
""")
    d.table(["Argument", "Meaning"], [
        ["--config FILE", "Criteria JSON file. Default: criterias.json in the current (startup) directory."],
        ["--lang CODE", "User interface language, e.g. hu or en (case-insensitive). Default: default_language of "
                        "locals.json (hu)."],
        ["--locals FILE", "Additional localization file merged over the built-in texts. Default: locals.json in "
                          "the current directory, if it exists."],
        ["IMAGE", "Optional image file opened at start (also works by dropping a file on the exe icon)."],
        ["-h, --help, /?", "Shows the description of the parameters in the selected language "
                           "(e.g. --lang en --help) and exits."],
        ["--version", "Shows the version window (name, version, copyright, source URL) and exits."],
    ], widths=[100, 370])
    d.h2("2.5 Command line checking and messages")
    d.p("The command line is parsed by app.parse_args (not argparse, so that every message can be localized). "
        "Values can be given as '--config FILE' or '--config=FILE', also with spaces around the equals sign "
        "('--config= FILE', '--config =FILE', '--config = FILE'); parameter names are case-insensitive, values "
        "are kept as given. app.process_command_line then checks it "
        "and returns a Startup object (run, help, version or error):")
    d.bullets([
        "Missing value (e.g. --config at the end, or followed by another parameter), a parameter given twice, "
        "an unknown parameter or more than one image file are errors.",
        "--config and --locals must name an existing file that can be loaded as a criteria file / localization "
        "file; an incompatible file (e.g. an image, or a localization file given as --config) is an error that "
        "includes the reason.",
        "--lang must be one of the available language codes (including languages from --locals).",
        "The IMAGE must exist and be a readable image.",
        "All errors are reported together, in the selected language, followed by a hint to use --help. The "
        "program then exits with exit code 2 without opening the window. Help and version exit with code 0.",
        "Problems with the implicit criterias.json / locals.json of the startup directory are not command line "
        "errors: they produce a warning in the window and the defaults are used.",
    ])
    d.p("Output: when the program has a console (run from source with python), help and version are printed to "
        "standard output and errors to standard error. The Windows exe is a GUI program without a console, so "
        "there the help is shown in a window (fixed-width font), the version in the same window as the About button (gui.AboutDialog) and the errors in message "
        "boxes.")

    d.h1("3. Architecture")
    d.p("The program is a Python package (trademarkimage) with a strict separation between the "
        "image logic, the workflow model and the user interface. Everything except gui.py can be used "
        "and tested without a display.")
    d.table(["Module", "Responsibility"], [
        ["criteria.py", "Criteria dataclass with defaults, validation, JSON load/save (load_criteria, save_criteria), "
                        "the list of image types and their Pillow modes."],
        ["analyzer.py", "analyze(path) -> ImageInfo (format, file size, displayed width/height, DPI, mode); "
                        "Prop enum of checked properties; check_value / check_image; value formatting."],
        ["converter.py", "Output formats and their capabilities (SAVE_FORMATS), prepare_image (orientation, "
                         "transparency, resize, type conversion) and convert_image (encoding with file size limit)."],
        ["session.py", "Session: GUI independent state of one examined image - the failing (flagged) properties, "
                       "the target values chosen by the user, input parsing, suggestions, keep-aspect logic and "
                       "building ConversionSettings."],
        ["gui.py", "Tkinter UI: App main window (with language selector), CriteriaPanel, ValueDialog; "
                   "drag and drop via tkinterdnd2."],
        ["i18n.py", "Localization: loads locals.json, merges extra files, keeps the current language, "
                    "tr(key, **values) returns a translated text with its placeholders filled in."],
        ["locals.json", "All user visible texts of every language (package data, also bundled into the exe)."],
        ["about.py", "Program identity: name and version, localized copyright line, source URL; used by the About button and --version."],
        ["app.py", "Command line parsing and checking, help/version/error output, localization set-up and "
                   "loading of the startup criteria; creates the window."],
        ["__main__.py, run_trademarkimage.py", "Entry points (python -m trademarkimage, PyInstaller)."],
    ], widths=[120, 350])
    d.h2("3.1 Data flow")
    d.steps([
        "app.main() parses the arguments and loads the criteria file (or the defaults).",
        "The user drops or opens an image; App.open_image() calls analyzer.analyze() and creates a Session.",
        "If Session.passed is true, a green OK button is shown; pressing it closes the program.",
        "Otherwise one red button per flagged property is shown. Pressing it opens a ValueDialog; the input is "
        "validated by Session.parse_value() and stored with Session.set_target(). Valid buttons turn white.",
        "When every property has a valid target value (Session.all_valid), App.convert_and_save() asks for the "
        "output file and format and calls converter.convert_image() with Session.conversion_settings().",
        "The written file is analysed again and checked against the criteria; the result is reported and the "
        "converted image is loaded into the window (which then shows the green OK button).",
    ])
    d.h2("3.2 Testability of the GUI")
    d.p("All dialogs used by App are wrapped in overridable methods (ask_value, ask_save_path, ask_yes_no, "
        "show_error, show_info, ask_criteria_save_path, ask_criteria_open_path). The GUI tests replace them "
        "with scripted answers and drive the real Tk widgets (button.invoke()).")

    d.h1("4. Criteria and the criteria file")
    d.table(["JSON key", "Default", "Meaning"], [
        ["max_file_size_mb", "2.0", "Maximum file size in MB (1 MB = 1 048 576 bytes)."],
        ["min_width", "945", "Minimum width in pixels."],
        ["min_height", "945", "Minimum height in pixels."],
        ["max_width", "2835", "Maximum width in pixels."],
        ["max_height", "2010", "Maximum height in pixels."],
        ["min_dpi", "96.0", "Minimum resolution in DPI (the smaller of the horizontal and vertical value is used)."],
        ["max_dpi", "600.0", "Maximum resolution in DPI (the larger of the horizontal and vertical value is used)."],
        ["allowed_types", '["RGB", "greyscale", "B&W", "CMYK"]',
         "Allowed image types. Known types: RGB, RGBA, greyscale, B&W, CMYK, palette."],
        ["allow_progressive_jpeg", "false", "If false, progressive JPEG files fail the check (baseline JPEG required). Other formats are not affected."],
    ], widths=[100, 130, 240])
    d.p("Example (criterias.example.json):")
    d.code("""
{
  "max_file_size_mb": 2.0,
  "min_width": 945,
  "min_height": 945,
  "max_width": 2835,
  "max_height": 2010,
  "min_dpi": 96.0,
  "max_dpi": 600.0,
  "allowed_types": ["RGB", "greyscale", "B&W", "CMYK"],
  "allow_progressive_jpeg": false
}
""")
    d.bullets([
        "Missing keys take their default value, so a file may contain only the values that differ.",
        "Unknown keys, wrong value types, non-integer pixel values, min > max, an empty or unknown type list "
        "are rejected with a descriptive error (CriteriaError).",
        "At start-up the program loads criterias.json from the current directory, or the file given with "
        "--config. A missing default file silently means default criteria; a missing --config file or an "
        "invalid file produces a warning and the defaults are used.",
        "When the user applies changed criteria in the GUI, the program asks whether to save them; the save "
        "dialog proposes the current criteria file. The 'Load criteria...' button loads an alternative file.",
        "Changing or loading criteria re-examines the current image.",
    ])

    d.h1("5. Checking rules")
    d.table(["Property", "Source", "Passes when"], [
        ["File size", "Size of the file on disk", "0 < size <= max_file_size_mb x 1 048 576"],
        ["Width / Height", "Pixel size after applying the EXIF orientation tag", "min <= value <= max"],
        ["Resolution", "dpi info (JPEG JFIF/EXIF, PNG pHYs, TIFF, BMP); JFIF dots/cm converted; "
                       "rounded to 0.1 DPI", "smaller of x/y DPI >= min_dpi and larger <= max_dpi; a missing value fails"],
        ["Image type", "Pillow mode mapped to a type name", "type is in allowed_types"],
        ["JPEG encoding", "JPEG files only: Pillow reports 'progressive' for a progressive (SOF2) file", "baseline, or allow_progressive_jpeg is true; other formats always pass"],
    ], widths=[85, 210, 175])
    d.table(["Pillow mode(s)", "Image type"], [
        ["1", "B&W"],
        ["L, LA, I, I;16, F", "greyscale"],
        ["RGB, RGBX, YCbCr", "RGB"],
        ["RGBA", "RGBA"],
        ["CMYK", "CMYK"],
        ["P, PA", "palette"],
        ["others (e.g. LAB, HSV)", "other - never allowed"],
    ], widths=[200, 270])

    d.h1("6. Workflow model (Session)")
    d.bullets([
        ("Flagged properties: ", "the properties that failed the initial check, plus any property that becomes "
                                 "invalid later (e.g. height after changing the width with 'keep aspect ratio')."),
        ("Target values: ", "initialised from the image; a button is red while its target is invalid and white "
                            "after a valid value is entered. Invalid input is rejected and the dialog stays open."),
        ("Keep aspect ratio: ", "setting width or height also sets the other dimension from the original ratio."),
        ("Suggestions: ", "the dialogs are pre-filled with a valid value; for the dimensions the scale factor "
                          "closest to 1 that brings both dimensions into range is used, so the default keeps "
                          "the aspect ratio when possible."),
        ("File size target: ", "only limits the output if file size was flagged; otherwise the criteria maximum "
                               "applies."),
        ("Automatic conversion: ", "when the last red button turns white the save dialog opens automatically; "
                                   "if cancelled, the 'Convert and save...' button can be used later."),
    ])

    d.h1("7. Conversion")
    d.steps([
        "The image is rotated according to the EXIF orientation tag.",
        "Transparency is composited onto a white background unless the target type is RGBA.",
        "The image is resized with the Lanczos filter in a continuous-tone working mode (L, RGB, RGBA or CMYK).",
        "The type is converted: B&W uses a fixed threshold of 128 (no dithering, keeps logos crisp); palette "
        "uses an adaptive 256 colour palette; CMYK uses Pillow's plain RGB->CMYK conversion (no ICC colour "
        "management).",
        "The image is encoded in memory with the chosen DPI (the original x/y values are kept when the resolution was not changed). JPEG is always written as baseline (not progressive) and uses the highest quality between 95 and 40 "
        "that fits under the size limit (binary search). If the file is still too big, the image is scaled down "
        "proportionally, but never below the minimum width/height; if that is not enough an error is shown.",
        "The data is written to the chosen file, then the file is analysed and checked again.",
    ])
    d.table(["Format", "Extensions", "Image types it can store", "Notes"], [
        ["JPEG", ".jpg .jpeg", "RGB, greyscale, CMYK", "Lossy; quality adjusted to the size limit"],
        ["PNG", ".png", "RGB, RGBA, greyscale, B&W, palette", "Lossless, optimize=True"],
        ["TIFF", ".tif .tiff", "all six types", "LZW, Group 4 for B&W"],
        ["BMP", ".bmp", "RGB, greyscale, B&W, palette", "Uncompressed - large files"],
    ], widths=[50, 75, 175, 170])
    d.p("The default output format is the format of the original file if it is writable and can store the "
        "target type; otherwise JPEG, PNG or TIFF (the first that fits). If the user selects an incompatible "
        "format, an error lists the suitable formats and the save dialog is shown again. A file name without "
        "extension gets the extension of the file type selected in the dialog.")

    d.h1("8. Localization")
    d.p("Every user visible text (window texts, buttons, dialogs, error messages, command line help) is "
        "stored in trademarkimage\\locals.json and looked up by a key with i18n.tr(key, **values). Texts may "
        "contain {placeholders} (Python str.format syntax, e.g. {value:.2f}); \\n in a text starts a new line.")
    d.code("""
{
  "default_language": "hu",
  "languages": {
    "hu": { "language_name": "Magyar",  "prop.width": "Szélesség", ... },
    "en": { "language_name": "English", "prop.width": "Width", ... }
  }
}
""")
    d.bullets([
        ("Default language: ", "Hungarian (default_language = hu). It can be changed with --lang, by setting "
                               "default_language in a locals.json in the startup directory, or at run time with "
                               "the Language (Nyelv) selector in the top right corner of the main window. "
                               "Switching rebuilds the window; the loaded image and the chosen target values "
                               "are kept."),
        ("Loading order: ", "the built-in file (inside the package / exe) is always loaded; then a locals.json in "
                            "the startup directory, or the file given with --locals, is merged over it key by "
                            "key. Such a file may contain only the changed texts or only a new language."),
        ("Fallback: ", "a text missing from the selected language is taken from English; an unknown key is "
                       "shown as the key itself. If an edited text has a wrong placeholder, the English text is "
                       "used instead, so a translation mistake cannot crash the program."),
        ("Image type names: ", "the JSON criteria file and the code use fixed type ids (RGB, greyscale, B&W ...); "
                               "only their display names (type.* keys) are translated."),
        ("Errors: ", "an invalid file given with --locals or an unknown --lang code is a command line error "
                     "(see 2.5); an invalid locals.json in the startup directory only produces a warning and the "
                     "built-in texts are used."),
    ])
    d.p("Adding a language: copy the \"en\" block in locals.json, change its code (e.g. \"de\"), set "
        "language_name and translate the values without changing the keys or placeholders. With the exe, edit the "
        "locals.json placed next to TrademarkImage.exe; no rebuild is needed. The test "
        "test_translation_is_complete checks that the built-in languages have the same keys and placeholders "
        "(add a new built-in language code to its parameter list).")

    d.h1("9. Tests")
    d.code(".venv\\Scripts\\python.exe -m pytest -q")
    d.p("162 automated tests (tests directory) cover:")
    d.bullets([
        "test_criteria.py - defaults, JSON round trip, partial files, every validation error.",
        "test_analyzer.py - mode classification, DPI reading and rounding, missing DPI, EXIF orientation, "
        "invalid files, boundary values of every criterion.",
        "test_converter.py - format helpers, resizing, DPI, conversion to every type and format, transparency "
        "flattening, JPEG quality reduction, down-scaling of lossless formats, impossible size limits.",
        "test_session.py - flagged properties, input parsing and validation, keep aspect ratio, suggestions, "
        "conversion settings.",
        "test_gui.py - drag and drop registration, green OK button and exit, red/white buttons, the full "
        "fix-and-convert workflow, incompatible format retry, criteria apply/save prompt, loading alternative "
        "criteria. These tests need a Windows desktop session (Tk) and are skipped when Tk is unavailable.",
        "test_i18n.py - Hungarian default, completeness of the translations (same keys and placeholders), every "
        "key used in the code exists, English fallback, merging extra files and adding a language, invalid "
        "files, translated messages, and in the GUI: Hungarian start-up, live language switch keeping the "
        "session, translated image type names in the value dialog.",
        "test_cli.py - parsing (both value forms, missing values, duplicates, unknown parameters, extra "
        "arguments), help in Hungarian and English, version, every kind of invalid --config / --locals / --lang / "
        "image value, errors in the selected language, defaults from the startup directory, exit codes, and "
        "output to the console or - without a console, as in the exe - to a window.",
    ])
    d.p("The tests run in English (an autouse fixture in conftest.py selects it) except the localization tests.")

    d.h1("10. Building the Windows executable")
    d.p("The build uses a separate, minimal virtual environment (.venv-build) so that development packages are "
        "not bundled. Run from the project directory:")
    d.code("powershell -ExecutionPolicy Bypass -File build_exe.ps1 [-Python <python.exe>] [-SkipTests]")
    d.steps([
        "Creates .venv-build with the given Python (default: python on PATH) if it does not exist.",
        "Installs requirements.txt, pytest and PyInstaller into it.",
        "Runs the test suite; the build stops if any test fails.",
        "Runs PyInstaller with TrademarkImage.spec: one-file, windowed (no console), tkinterdnd2 data and the "
        "native tkdnd library (Windows builds only) collected, locals.json bundled as package data, "
        "development packages excluded.",
        "Result: dist\\TrademarkImage.exe (about 19 MB) plus dist\\criterias.example.json and an editable "
        "dist\\locals.json.",
    ])
    d.p("Deployment: copy TrademarkImage.exe (and optionally locals.json) to the target machine. To use custom "
        "criteria, place a criterias.json next to the exe (when started by double click the startup directory is "
        "the exe folder) or start it with --config. The exe works without locals.json, using its built-in texts. "
        "The one-file exe unpacks itself into a temporary folder at start, so the first start takes a few "
        "seconds. Unsigned executables may trigger a Windows SmartScreen warning.")

    d.h1("11. Known limitations")
    d.bullets([
        "Only the first frame of multi-frame images (animated GIF, multi-page TIFF) is processed.",
        "CMYK conversion is not colour managed; embedded ICC profiles are not carried over.",
        "Metadata (EXIF, ICC, XMP) is not copied to the converted image; only the DPI is written.",
        "Changing the DPI only changes the stored metadata; pixel dimensions are controlled by width/height.",
        "Numbers are always shown with a decimal point; input accepts both decimal point and comma.",
        "The language chosen in the window is not remembered; use --lang or default_language in locals.json "
        "to change the start-up language permanently.",
    ])

    d.h1("12. Regenerating the documentation")
    d.p("This technical documentation and the user manuals (English and Hungarian) are generated by "
        "docs\\make_docs.py from a single source, in Word (.docx) and "
        "Markdown (.md) format, so the two formats always have the same content. Edit the texts there, "
        "not in the generated files.")
    d.code(".venv\\Scripts\\python.exe docs\\make_docs.py")
    d.save("TrademarkImage_Technical_Documentation_EN")


LOCALS_STRUCTURE = """
{
  "default_language": "hu",
  "languages": {
    "hu": { "language_name": "Magyar",  "prop.width": "Szélesség", ... },
    "en": { "language_name": "English", "prop.width": "Width", ... }
  }
}
"""

CRITERIA_EXAMPLE = """
{
  "max_file_size_mb": 2.0,
  "min_width": 945,
  "min_height": 945,
  "max_width": 2835,
  "max_height": 2010,
  "min_dpi": 96.0,
  "max_dpi": 600.0,
  "allowed_types": ["RGB", "greyscale", "B&W", "CMYK"],
  "allow_progressive_jpeg": false
}
"""


def user_guide_hu(Doc):
    d = Doc("Védjegykép-ellenőrző", "Felhasználói kézikönyv (Trademark Image Checker, 1.2.0 verzió)")
    d.p("A kézikönyv két részből áll. Az I. rész a program mindennapi használatát írja le. A II. rész a "
        "haladó lehetőségeket mutatja be: a parancssori paramétereket, a nyelv beállítását és a nyelvi fájlt, "
        "valamint a feltételfájl felépítését.")

    # ----- I. rész -----
    d.h1("I. Alapvető használat")

    d.h2("1. Mire való a program?")
    d.p("A Védjegykép-ellenőrző egy képfájlról megállapítja, hogy megfelel-e a beállított "
        "feltételeknek (fájlméret, képméret pixelben, felbontás, képtípus), például egy védjegybejelentéshez "
        "feltöltendő kép követelményeinek. Ha a kép nem felel meg, a program segít beállítani a megfelelő "
        "értékeket, majd elkészíti és elmenti az átalakított képet.")

    d.h2("2. A program indítása")
    d.bullets([
        "Kattintson duplán a TrademarkImage.exe fájlra. Python telepítésére nincs szükség.",
        "Az első indítás néhány másodpercig tarthat, mert a program kicsomagolja magát.",
        "Ha a Windows SmartScreen figyelmeztet, válassza a „További információ”, majd a „Futtatás mindenképp” "
        "lehetőséget (a program nincs digitálisan aláírva).",
        "Egy képfájlt az exe ikonjára húzva a program azonnal azzal a képpel indul.",
    ])
    d.p("A program magyarul indul. A nyelv a főablak jobb felső sarkában, a Nyelv listában bármikor "
        "átváltható (Magyar / English); a betöltött kép és a beállított értékek ilyenkor megmaradnak. A nyelv "
        "tartós beállítását a 11. fejezet írja le.")

    d.h2("3. A főablak részei")
    d.table(["Terület", "Leírás"], [
        ["Névjegy (bal felső sarok)", "Megjeleníti a program verziószámát, a szerzői jogot és a forráskód címét (a cím kattintható)."],
        ["Nyelv (jobb felső sarok)", "A felület nyelvének kiválasztása (Magyar / English)."],
        ["Feltételek (bal oldal)", "A feltételek beállítása, valamint az Alkalmaz, a Feltételek betöltése... és a "
                                   "Feltételek mentése másként... gomb. Alul látható, melyik fájlból "
                                   "származnak a feltételek."],
        ["Képterület (jobb felül)", "Ide kell behúzni a képet; betöltés után a kép előnézete látszik."],
        ["Képadatok", "Fájl, formátum, mód, méret pixelben, felbontás (DPI), fájlméret."],
        ["Eredmény (jobb alul)", "Zöld OK gomb, vagy a nem teljesülő feltételek piros gombjai és az "
                                 "Átalakítás és mentés... gomb."],
    ], widths=[130, 340])

    d.h2("4. Kép betöltése")
    d.steps([
        "Húzza a képfájlt a Windows Intézőből a „Húzzon ide egy képet” feliratú területre, vagy",
        "kattintson a területre, és válassza ki a képet a megnyíló ablakban.",
    ])
    d.p("Olvasható formátumok: JPEG, PNG, TIFF, BMP, GIF, WebP és a Pillow könyvtár által ismert egyéb "
        "formátumok. Ha a fájl nem kép, hibaüzenet jelenik meg.")

    d.h2("5. Az ellenőrzés eredménye")
    d.h3("5.1 A kép megfelel")
    d.p("Ha minden feltétel teljesül, egy zöld gomb jelenik meg: „OK - a kép minden feltételnek megfelel "
        "(kattintson a kilépéshez)”. Erre kattintva a program bezárul, a képen nem változtat semmit.")
    d.h3("5.2 A kép nem felel meg")
    d.p("Minden nem teljesülő feltételhez egy piros gomb tartozik, amelyen a feltétel neve, a kép jelenlegi "
        "értéke és zárójelben a követelmény olvasható, például: „Szélesség: 4000 px (945-2835 px)”. "
        "A lehetséges gombok: Fájlméret, Szélesség, Magasság, Felbontás, Képtípus és JPEG kódolás. A felbontásnál "
        "„nincs megadva” látható, ha a fájl nem tárol DPI értéket.")

    d.h2("6. A hibás értékek javítása")
    d.steps([
        "Kattintson egy piros gombra. Megnyílik egy ablak, amely mutatja a jelenlegi értéket és a "
        "követelményt, az Új érték mező pedig egy megfelelő javaslattal van kitöltve.",
        "Írja be (vagy fogadja el) a kívánt értéket, és nyomja meg az OK gombot.",
        "Ha az érték megfelel a feltételnek, a gomb fehér lesz, és az új értéket mutatja (pl. „Szélesség: "
        "4000 px -> 2800 px”). Ha nem felel meg, hibaüzenet jelenik meg, és az ablak nyitva marad. A Mégse "
        "gombbal a változtatás elvethető, a gomb piros marad.",
    ])
    d.bullets([
        ("Szélesség, magasság: ", "a „Méretarány megtartása” jelölőnégyzet bekapcsolva (alapértelmezés) a "
                                  "másik méretet is arányosan módosítja, így a kép nem torzul. Ha emiatt a "
                                  "másik méret kerül a megengedett tartományon kívülre, arra is megjelenik egy "
                                  "piros gomb. A felajánlott értékek - ha lehetséges - megtartják a képarányt."),
        ("Felbontás: ", "egész szám DPI-ben, a legkisebb és a legnagyobb felbontás között (alapesetben 96-600 DPI). A felbontás csak a fájlban tárolt adat, a pixelméretet nem "
                        "változtatja."),
        ("Fájlméret: ", "a megengedett legnagyobb méret MB-ban (tizedesvessző és -pont is használható). A "
                        "program a mentéskor szükség esetén csökkenti a JPEG minőséget, illetve a képméretet, "
                        "hogy a fájl beleférjen."),
        ("Képtípus: ", "a listából választható ki a megengedett típusok egyike."),
        ("JPEG kódolás: ", "progresszív JPEG esetén válassza az „alap (baseline)” értéket; a program a képet alap JPEG-ként menti újra."),
    ])

    d.h2("7. Átalakítás és mentés")
    d.p("Amikor az utolsó piros gomb is fehér lesz, a program automatikusan megnyitja a mentési ablakot. Ha ezt "
        "megszakítja, később az „Átalakítás és mentés...” gombbal indíthatja újra.")
    d.bullets([
        "Az alapértelmezett formátum az eredeti fájl formátuma (ha az a választott képtípust tárolni tudja), "
        "az alapértelmezett fájlnév „<eredeti név>_atalakitott”.",
        "A formátumot a fájltípus-listából vagy a kiterjesztéssel (.jpg, .png, .tif, .bmp) lehet megadni.",
        "Ha a formátum nem tudja tárolni a képtípust (pl. CMYK képet PNG-be), a program jelzi, mely formátumok "
        "megfelelőek, és újra megnyitja a mentési ablakot.",
        "Mentés után a program újra ellenőrzi az elkészült fájlt, kiírja az eredményt (méret, fájlméret, "
        "JPEG minőség), és betölti az új képet - ha minden rendben, megjelenik a zöld OK gomb.",
    ])
    d.table(["Formátum", "Tárolható képtípusok", "Megjegyzés"], [
        ["JPEG (.jpg)", "RGB, szürkeárnyalatos, CMYK", "Veszteséges, kis fájlméret - általában ez ajánlott"],
        ["PNG (.png)", "RGB, RGBA, szürkeárnyalatos, fekete-fehér, palettás", "Veszteségmentes"],
        ["TIFF (.tif)", "mind", "Veszteségmentes tömörítés"],
        ["BMP (.bmp)", "RGB, szürkeárnyalatos, fekete-fehér, palettás", "Tömörítetlen, nagy fájlok"],
    ], widths=[85, 175, 210])
    d.p("Átalakításkor a JPEG fájlok mindig alap (nem progresszív) kódolással készülnek, az átlátszó részek fehér hátteret kapnak (kivéve RGBA típus esetén), a fekete-fehér "
        "átalakítás 50%-os küszöbbel történik, a kép pedig a fényképezőgép által tárolt tájolás szerint "
        "elforgatásra kerül.")

    d.h2("8. A feltételek beállítása")
    d.table(["Mező", "Alapérték"], [
        ["Legnagyobb fájlméret (MB)", "2 (1 MB = 1 048 576 bájt)"],
        ["Legkisebb szélesség (px)", "945"],
        ["Legkisebb magasság (px)", "945"],
        ["Legnagyobb szélesség (px)", "2835"],
        ["Legnagyobb magasság (px)", "2010"],
        ["Legkisebb felbontás (DPI)", "96"],
        ["Legnagyobb felbontás (DPI)", "600"],
        ["Megengedett képtípusok", "RGB, szürkeárnyalatos, fekete-fehér, CMYK"],
        ["Progresszív JPEG megengedett", "nem (kikapcsolva)"],
    ], widths=[200, 270])
    d.table(["Képtípus", "Jelentése"], [
        ["RGB (színes)", "Színes kép"],
        ["RGBA (színes, átlátszó)", "Színes kép átlátszósággal"],
        ["szürkeárnyalatos", "Szürkeárnyalatos kép"],
        ["fekete-fehér", "Kétszínű (fekete-fehér) kép"],
        ["CMYK", "Nyomdai színmodellű kép"],
        ["palettás", "Legfeljebb 256 színű kép, pl. GIF"],
    ], widths=[200, 270])
    d.steps([
        "Módosítsa a mezőket, majd nyomja meg az Alkalmaz gombot.",
        "Ha a feltételek változtak, a program megkérdezi, el akarja-e menteni őket fájlba. Igen válasz esetén "
        "a mentési ablak a jelenlegi feltételfájlt (alapesetben criterias.json) ajánlja fel.",
        "A betöltött kép ellenőrzése az új feltételekkel automatikusan megismétlődik.",
    ])
    d.p("Más feltételkészlet betöltése: Feltételek betöltése... gomb, majd a JSON fájl kiválasztása. "
        "Mentés tetszőleges néven: Feltételek mentése másként... gomb. A feltételfájlokról bővebben a 12. "
        "fejezet szól.")

    d.h2("9. Gyakori kérdések és hibaüzenetek")
    d.table(["Üzenet / helyzet", "Teendő"], [
        ["A kép nem nyitható meg / Nem felismerhető képfájl", "A fájl nem kép vagy sérült - válasszon másik "
                                                              "fájlt."],
        ["Érvénytelen érték", "A beírt érték nem szám, vagy kívül esik a megengedett tartományon."],
        ["Érvénytelen feltételek", "A Feltételek panel valamelyik mezője hibás (pl. a minimum nagyobb a "
                                   "maximumnál, vagy nincs megengedett képtípus)."],
        ["Nem megfelelő formátum", "A választott fájlformátum nem tudja tárolni a képtípust - válasszon a "
                                   "felsorolt formátumok közül."],
        ["Sikertelen átalakítás ... legkisebb megengedett méret", "A kép a megadott fájlméretbe a minimális "
                                                                   "képméret mellett sem fér bele. Válasszon "
                                                                   "JPEG formátumot vagy más képtípust."],
        ["A felbontás „nincs megadva”", "A fájl nem tárol DPI értéket; állítson be egyet a Felbontás gombbal."],
    ], widths=[170, 300])

    # ----- II. rész -----
    d.page_break()
    d.h1("II. Haladó használat és testreszabás")

    d.h2("10. Parancssori paraméterek")
    d.p("A program parancssorból (vagy parancsikonból) paraméterekkel is indítható:")
    d.code("""
TrademarkImage.exe
TrademarkImage.exe kep.jpg
TrademarkImage.exe --config sajat_feltetelek.json kep.jpg
TrademarkImage.exe --lang en
TrademarkImage.exe --locals sajat_szovegek.json
""")
    d.table(["Paraméter", "Jelentése"], [
        ["--config FÁJL", "A feltételeket tartalmazó JSON fájl (alapértelmezés: criterias.json az indítási "
                          "könyvtárban). Lásd a 12. fejezetet."],
        ["--lang KÓD", "A felület nyelve: hu (magyar) vagy en (angol), illetve bármely további, a locals.json-ba "
                       "felvett nyelv kódja. Lásd a 11. fejezetet."],
        ["--locals FÁJL", "Saját szövegfájl (alapértelmezés: locals.json az indítási könyvtárban, ha van)."],
        ["KÉP", "Induláskor megnyitandó képfájl."],
        ["-h, --help", "Megjeleníti a paraméterek leírását. Angolul: --lang en --help."],
        ["--version", "Megjeleníti a verzióablakot (ugyanazt, mint a Névjegy gomb)."],
    ], widths=[110, 360])
    d.bullets([
        "Az érték szóközzel vagy egyenlőségjellel adható meg; az egyenlőségjel körül szóköz is lehet: "
        "--lang en, --lang=en, --lang = en. A szóközt tartalmazó fájlneveket tegye idézőjelek közé.",
        "A paraméterek nevében a kis- és nagybetű nem számít (--LANG=en is jó).",
        "Az indítási könyvtár dupla kattintásos indításnál az exe mappája. Parancsikon esetén a parancsikon "
        "„Indítás helye” mezője adja meg.",
        "A súgó egy külön ablakban, a verziószám a Névjegy ablakban jelenik meg; ezután a program kilép.",
    ])
    d.h3("10.1 Hibás paraméterek")
    d.p("A program induláskor ellenőrzi a paramétereket. Ha valamelyik hibás, egy hibaüzenetben felsorolja az "
        "összes problémát, és nem indul el. Hibának számít például:")
    d.bullets([
        "a --config, a --locals vagy a --lang után hiányzik az érték;",
        "a megadott fájl nem létezik, mappa, vagy nem megfelelő (pl. kép vagy szövegfájl a --config után);",
        "ismeretlen nyelvkód a --lang után (az üzenet felsorolja az elérhető nyelveket);",
        "a megadott kép nem létezik vagy nem kép;",
        "ismeretlen vagy többször megadott paraméter, illetve egynél több képfájl.",
    ])
    d.p("A hibaüzenet a kiválasztott nyelven jelenik meg, és a végén megmutatja, hogyan kérhető súgó. "
        "Szkriptekből a kilépési kód is vizsgálható: 0 = rendben, 2 = parancssori hiba.")

    d.h2("11. A nyelv beállítása és a nyelvi fájl")
    d.h3("11.1 A nyelv kiválasztása")
    d.bullets([
        ("Futás közben: ", "a főablak jobb felső sarkában lévő Nyelv listában. A választás csak az adott "
                           "futásra érvényes."),
        ("Indításkor: ", "a --lang paraméterrel, pl. TrademarkImage.exe --lang en."),
        ("Tartósan: ", "a locals.json fájlban a \"default_language\" értékének átírásával (pl. \"en\")."),
    ])
    d.h3("11.2 A locals.json felépítése")
    d.p("A program összes szövege a locals.json fájlban van, amely az exe mellett található, és "
        "Jegyzettömbbel szerkeszthető (UTF-8 kódolással mentse). Felépítése:")
    d.code(LOCALS_STRUCTURE)
    d.bullets([
        "Egy szöveg módosításához írja át az idézőjelek közötti értéket a kettőspont után. A kulcsot (a kettőspont "
        "előtti részt) ne változtassa meg.",
        "A kapcsos zárójelben lévő részeket (pl. {value}, {path}) a program tölti ki - ezeket változatlanul hagyja "
        "meg. A \\n új sort jelent.",
        "Ha egy szöveg hiányzik egy nyelvből, a program az angol szöveget mutatja. Hibás fájl esetén "
        "figyelmeztetést ad, és a beépített szövegekkel dolgozik.",
        "Elég csak a megváltoztatott szövegeket tartalmazó fájlt az exe mellé tenni: a program a beépített "
        "szövegekhez hozzáadja, illetve azokat felülírja vele.",
    ])
    d.h3("11.3 Új nyelv felvétele")
    d.steps([
        "Nyissa meg a locals.json fájlt, és másolja le az \"en\" blokkot.",
        "Adjon neki új kódot (pl. \"de\"), és írja át a \"language_name\" értékét (pl. \"Deutsch\").",
        "Fordítsa le a szövegeket, majd mentse a fájlt UTF-8 kódolással.",
        "A program következő indításakor az új nyelv megjelenik a Nyelv listában.",
    ])

    d.h2("12. A feltételfájl")
    d.p("Induláskor a program az indítási könyvtárban lévő criterias.json fájlból tölti be a feltételeket, vagy "
        "a --config paraméterrel megadott fájlból. Ha azt szeretné, hogy a program mindig a saját feltételeivel "
        "induljon, mentse őket criterias.json néven az exe mellé (Feltételek mentése másként... gomb). "
        "A feltételfájl egyszerű JSON szöveg, Jegyzettömbbel is szerkeszthető:")
    d.code(CRITERIA_EXAMPLE)
    d.table(["Kulcs", "Mező a felületen", "Alapérték"], [
        ["max_file_size_mb", "Legnagyobb fájlméret (MB)", "2.0"],
        ["min_width", "Legkisebb szélesség (px)", "945"],
        ["min_height", "Legkisebb magasság (px)", "945"],
        ["max_width", "Legnagyobb szélesség (px)", "2835"],
        ["max_height", "Legnagyobb magasság (px)", "2010"],
        ["min_dpi", "Legkisebb felbontás (DPI)", "96.0"],
        ["max_dpi", "Legnagyobb felbontás (DPI)", "600.0"],
        ["allowed_types", "Megengedett képtípusok", "[\"RGB\", \"greyscale\", \"B&W\", \"CMYK\"]"],
        ["allow_progressive_jpeg", "Progresszív JPEG megengedett", "false (true = megengedett)"],
    ], widths=[120, 180, 170])
    d.p("A képtípusokat a fájlban az azonosítójukkal kell megadni:")
    d.table(["Azonosító a fájlban", "Képtípus a felületen"], [
        ["RGB", "RGB (színes)"],
        ["RGBA", "RGBA (színes, átlátszó)"],
        ["greyscale", "szürkeárnyalatos"],
        ["B&W", "fekete-fehér"],
        ["CMYK", "CMYK"],
        ["palette", "palettás"],
    ], widths=[200, 270])
    d.bullets([
        "A hiányzó kulcsok az alapértéket kapják, így elég csak az eltérő értékeket megadni.",
        "Tizedesponttal kell írni (2.5), a pixelértékek egész számok.",
        "Ismeretlen kulcs, hibás érték vagy ellentmondás (pl. a minimum nagyobb a maximumnál) esetén: ha a "
        "fájlt a --config adta meg, a program hibaüzenettel kilép; az indítási könyvtár criterias.json fájljánál "
        "figyelmeztet, és az alapértelmezett feltételekkel indul.",
    ])

    d.h2("13. Hibaelhárítás a haladó beállításoknál")
    d.table(["Üzenet / helyzet", "Teendő"], [
        ["Parancssori hiba, a program nem indul", "Javítsa a hibaüzenetben felsorolt paramétereket (lásd 10.1). "
                                                  "A paraméterek leírása: --help."],
        ["Figyelmeztetés indításkor", "Az indítási könyvtárban lévő criterias.json vagy locals.json hibás. A "
                                      "program az alapértelmezett beállításokkal fut."],
        ["Angol szöveg jelenik meg magyar felületen", "A szöveg hiányzik a locals.json magyar részéből - "
                                                      "pótolja a hiányzó kulcsot."],
        ["Kulcs (pl. prop.width) látszik szöveg helyett", "A szöveg az angol részből is hiányzik - pótolja."],
        ["Ékezetes betűk helyett furcsa jelek", "Mentse a locals.json-t UTF-8 kódolással."],
    ], widths=[170, 300])
    d.save("TrademarkImage_Felhasznaloi_kezikonyv_HU")


def user_guide_en(Doc):
    d = Doc("Trademark Image Checker", "User manual (version 1.2.0)")
    d.p("This manual has two parts. Part I describes the everyday use of the program. Part II covers the "
        "advanced options: command line parameters, the language setting and the language file, and the "
        "structure of the criteria file.")

    # ----- Part I -----
    d.h1("I. Basic usage")

    d.h2("1. What is the program for?")
    d.p("Trademark Image Checker examines an image file and tells whether it meets the configured criteria "
        "(file size, image size in pixels, resolution, image type), for example the requirements for an image "
        "uploaded with a trademark application. If the image does not meet them, the program helps you set "
        "suitable values, then creates and saves the converted image.")

    d.h2("2. Starting the program")
    d.bullets([
        "Double-click TrademarkImage.exe. Python does not need to be installed.",
        "The first start may take a few seconds because the program unpacks itself.",
        "If Windows SmartScreen shows a warning, choose \"More info\" and then \"Run anyway\" (the program is "
        "not digitally signed).",
        "If you drag an image file onto the exe icon, the program starts with that image.",
    ])
    d.p("The program starts in Hungarian by default. To switch to English, select English in the Nyelv list in "
        "the top right corner of the window; the loaded image and the values already set are kept. To start in "
        "English every time, see section 11.", bold_prefix="Note: ")

    d.h2("3. The main window")
    d.table(["Area", "Description"], [
        ["About (top left corner)", "Shows the version number, the copyright and the address of the source code (the address can be clicked)."],
        ["Language (top right corner)", "Selects the language of the user interface (Magyar / English)."],
        ["Criteria (left side)", "Setting the criteria, with the Apply, Load criteria... and Save criteria as... "
                                 "buttons. The bottom line shows which file the criteria come from."],
        ["Image area (top right)", "Drop the image here; after loading, a preview of the image is shown."],
        ["Image data", "File, format, mode, size in pixels, resolution (DPI), file size."],
        ["Result (bottom right)", "A green OK button, or red buttons for the criteria that are not met and the "
                                  "Convert and save... button."],
    ], widths=[130, 340])

    d.h2("4. Loading an image")
    d.steps([
        "Drag the image file from Windows Explorer onto the area labelled \"Drag and drop an image here\", or",
        "click the area and select the image in the dialog that opens.",
    ])
    d.p("Readable formats: JPEG, PNG, TIFF, BMP, GIF, WebP and other formats known by the Pillow library. If the "
        "file is not an image, an error message is shown.")

    d.h2("5. The result of the check")
    d.h3("5.1 The image meets the criteria")
    d.p("If every criterion is met, a green button appears: \"OK - the image meets all criteria (click to "
        "exit)\". Clicking it closes the program without changing the image.")
    d.h3("5.2 The image does not meet the criteria")
    d.p("There is a red button for every criterion that is not met. It shows the name of the criterion, the "
        "current value of the image and the requirement in brackets, for example: \"Width: 4000 px (945-2835 "
        "px)\". The possible buttons are File size, Width, Height, Resolution, Image type and JPEG encoding. For the resolution, "
        "\"not set\" is shown if the file does not store a DPI value.")

    d.h2("6. Fixing the values")
    d.steps([
        "Click a red button. A dialog opens showing the current value and the requirement; the New value field "
        "is filled in with a suitable suggestion.",
        "Type (or accept) the desired value and press OK.",
        "If the value meets the criterion, the button turns white and shows the new value (e.g. \"Width: "
        "4000 px -> 2800 px\"). If not, an error message is shown and the dialog stays open. Cancel discards the "
        "change and the button stays red.",
    ])
    d.bullets([
        ("Width, height: ", "when the \"Keep aspect ratio\" check box is on (default), the other dimension is "
                            "changed proportionally so the image is not distorted. If this moves the other "
                            "dimension out of the allowed range, a red button appears for it too. The suggested "
                            "values keep the aspect ratio whenever possible."),
        ("Resolution: ", "a whole number in DPI between the minimum and maximum resolution (96-600 DPI by default). The resolution is only information stored in the file; it does "
                         "not change the pixel size."),
        ("File size: ", "the largest allowed size in MB (both decimal point and comma are accepted). When "
                        "saving, the program lowers the JPEG quality or the image size if needed, so that the file "
                        "fits."),
        ("Image type: ", "choose one of the allowed types from the list."),
        ("JPEG encoding: ", "for a progressive JPEG choose \"baseline\"; the program saves the image again as a baseline JPEG."),
    ])

    d.h2("7. Converting and saving")
    d.p("When the last red button turns white, the program opens the save dialog automatically. If you cancel "
        "it, you can start it again later with the \"Convert and save...\" button.")
    d.bullets([
        "The default format is the format of the original file (if it can store the selected image type); the "
        "default file name is \"<original name>_converted\" (\"_atalakitott\" in Hungarian).",
        "The format can be chosen in the file type list or with the extension (.jpg, .png, .tif, .bmp).",
        "If the format cannot store the image type (e.g. a CMYK image as PNG), the program lists the suitable "
        "formats and opens the save dialog again.",
        "After saving, the program checks the new file again, shows the result (size, file size, JPEG quality) "
        "and loads the new image - if everything is fine, the green OK button appears.",
    ])
    d.table(["Format", "Image types it can store", "Note"], [
        ["JPEG (.jpg)", "RGB, greyscale, CMYK", "Lossy, small files - usually recommended"],
        ["PNG (.png)", "RGB, RGBA, greyscale, B&W, palette", "Lossless"],
        ["TIFF (.tif)", "all", "Lossless compression"],
        ["BMP (.bmp)", "RGB, greyscale, B&W, palette", "Uncompressed, large files"],
    ], widths=[85, 175, 210])
    d.p("During conversion JPEG files are always written as baseline (not progressive), transparent areas get a white background (except for the RGBA type), the black and "
        "white conversion uses a 50% threshold, and the image is rotated according to the orientation stored by "
        "the camera.")

    d.h2("8. Setting the criteria")
    d.table(["Field", "Default"], [
        ["Maximum file size (MB)", "2 (1 MB = 1 048 576 bytes)"],
        ["Minimum width (px)", "945"],
        ["Minimum height (px)", "945"],
        ["Maximum width (px)", "2835"],
        ["Maximum height (px)", "2010"],
        ["Minimum resolution (DPI)", "96"],
        ["Maximum resolution (DPI)", "600"],
        ["Allowed image types", "RGB, greyscale, B&W, CMYK"],
        ["Progressive JPEG allowed", "no (unchecked)"],
    ], widths=[200, 270])
    d.table(["Image type", "Meaning"], [
        ["RGB", "Colour image"],
        ["RGBA", "Colour image with transparency"],
        ["greyscale", "Greyscale image"],
        ["B&W", "Two-colour (black and white) image"],
        ["CMYK", "Image in the printing colour model"],
        ["palette", "Image with at most 256 colours, e.g. GIF"],
    ], widths=[200, 270])
    d.steps([
        "Change the fields, then press Apply.",
        "If the criteria have changed, the program asks whether you want to save them to a file. If you answer "
        "Yes, the save dialog offers the current criteria file (criterias.json by default).",
        "The loaded image is checked again with the new criteria automatically.",
    ])
    d.p("To load another set of criteria, press Load criteria... and select the JSON file. To save under any "
        "name, press Save criteria as.... Criteria files are described in section 12.")

    d.h2("9. Frequently asked questions and error messages")
    d.table(["Message / situation", "What to do"], [
        ["Cannot open image / Not a recognised image file", "The file is not an image or it is damaged - choose "
                                                            "another file."],
        ["Invalid value", "The value entered is not a number or is outside the allowed range."],
        ["Invalid criteria", "A field of the Criteria panel is wrong (e.g. the minimum is greater than the "
                             "maximum, or no image type is allowed)."],
        ["Unsuitable format", "The selected file format cannot store the image type - choose one of the listed "
                              "formats."],
        ["Conversion failed ... minimum size", "The image does not fit into the given file size even at the "
                                               "minimum image size. Choose the JPEG format or another image type."],
        ["Resolution is \"not set\"", "The file does not store a DPI value; set one with the Resolution button."],
        ["The window is in Hungarian", "Select English in the Nyelv list in the top right corner."],
    ], widths=[170, 300])

    # ----- Part II -----
    d.page_break()
    d.h1("II. Advanced usage and customization")

    d.h2("10. Command line parameters")
    d.p("The program can also be started from the command line (or from a shortcut) with parameters:")
    d.code("""
TrademarkImage.exe
TrademarkImage.exe image.jpg
TrademarkImage.exe --config my_criteria.json image.jpg
TrademarkImage.exe --lang en
TrademarkImage.exe --locals my_texts.json
""")
    d.table(["Parameter", "Meaning"], [
        ["--config FILE", "JSON file containing the criteria (default: criterias.json in the startup folder). "
                          "See section 12."],
        ["--lang CODE", "Language of the user interface: hu (Hungarian) or en (English), or the code of any "
                        "other language added to locals.json. See section 11."],
        ["--locals FILE", "Your own text file (default: locals.json in the startup folder, if present)."],
        ["IMAGE", "Image file to open at start."],
        ["-h, --help", "Shows the description of the parameters. In English: --lang en --help."],
        ["--version", "Shows the version window (the same as the About button)."],
    ], widths=[110, 360])
    d.bullets([
        "A value can be given after a space or an equals sign, with or without spaces around it: --lang en, "
        "--lang=en, --lang = en. Put file names containing spaces in quotation marks.",
        "Parameter names are not case-sensitive (--LANG=en works too).",
        "When the program is started by double-click, the startup folder is the folder of the exe. For a "
        "shortcut it is the \"Start in\" field of the shortcut.",
        "The help is shown in a separate window and the version number in the About window; then the program "
        "exits.",
    ])
    d.h3("10.1 Wrong parameters")
    d.p("The program checks the parameters at start. If any of them is wrong, it lists all problems in one "
        "error message and does not start. For example, these are errors:")
    d.bullets([
        "the value is missing after --config, --locals or --lang;",
        "the given file does not exist, is a folder, or is not suitable (e.g. an image or a text file after "
        "--config);",
        "an unknown language code after --lang (the message lists the available languages);",
        "the given image does not exist or is not an image;",
        "an unknown parameter or one given more than once, or more than one image file.",
    ])
    d.p("The error message is shown in the selected language and ends with a hint on how to get help. Scripts "
        "can also check the exit code: 0 = OK, 2 = command line error.")

    d.h2("11. Language setting and the language file")
    d.h3("11.1 Choosing the language")
    d.bullets([
        ("While running: ", "in the Language (Nyelv) list in the top right corner of the main window. The choice "
                            "is valid for the current run only."),
        ("At start: ", "with the --lang parameter, e.g. TrademarkImage.exe --lang en."),
        ("Permanently: ", "by changing the \"default_language\" value in locals.json (e.g. to \"en\")."),
    ])
    d.h3("11.2 Structure of locals.json")
    d.p("All texts of the program are in the locals.json file next to the exe. It can be edited with Notepad "
        "(save it with UTF-8 encoding). Its structure:")
    d.code(LOCALS_STRUCTURE)
    d.bullets([
        "To change a text, edit the value in quotation marks after the colon. Do not change the key (the part "
        "before the colon).",
        "Parts in curly brackets (e.g. {value}, {path}) are filled in by the program - leave them unchanged. "
        "\\n starts a new line.",
        "If a text is missing from a language, the program shows the English text. If the file is invalid, the "
        "program shows a warning and uses its built-in texts.",
        "It is enough to put a file containing only the changed texts next to the exe: the program adds it to "
        "its built-in texts and overrides them with it.",
    ])
    d.h3("11.3 Adding a language")
    d.steps([
        "Open locals.json and copy the \"en\" block.",
        "Give it a new code (e.g. \"de\") and change \"language_name\" (e.g. \"Deutsch\").",
        "Translate the texts, then save the file with UTF-8 encoding.",
        "The new language appears in the Language list the next time the program starts.",
    ])

    d.h2("12. The criteria file")
    d.p("At start the program loads the criteria from criterias.json in the startup folder, or from the file "
        "given with the --config parameter. If you want the program to always start with your own criteria, "
        "save them as criterias.json next to the exe (Save criteria as... button). The criteria file is plain "
        "JSON text and can also be edited with Notepad:")
    d.code(CRITERIA_EXAMPLE)
    d.table(["Key", "Field in the window", "Default"], [
        ["max_file_size_mb", "Maximum file size (MB)", "2.0"],
        ["min_width", "Minimum width (px)", "945"],
        ["min_height", "Minimum height (px)", "945"],
        ["max_width", "Maximum width (px)", "2835"],
        ["max_height", "Maximum height (px)", "2010"],
        ["min_dpi", "Minimum resolution (DPI)", "96.0"],
        ["max_dpi", "Maximum resolution (DPI)", "600.0"],
        ["allowed_types", "Allowed image types", "[\"RGB\", \"greyscale\", \"B&W\", \"CMYK\"]"],
        ["allow_progressive_jpeg", "Progressive JPEG allowed", "false (true = allowed)"],
    ], widths=[120, 180, 170])
    d.p("Image types must be given by their id in the file:")
    d.table(["Id in the file", "Image type in the window (English / Hungarian)"], [
        ["RGB", "RGB / RGB (színes)"],
        ["RGBA", "RGBA / RGBA (színes, átlátszó)"],
        ["greyscale", "greyscale / szürkeárnyalatos"],
        ["B&W", "B&W / fekete-fehér"],
        ["CMYK", "CMYK / CMYK"],
        ["palette", "palette / palettás"],
    ], widths=[200, 270])
    d.bullets([
        "Missing keys get their default value, so it is enough to list the values that differ.",
        "Use a decimal point (2.5); pixel values are whole numbers.",
        "If there is an unknown key, a wrong value or a contradiction (e.g. the minimum is greater than the "
        "maximum): for a file given with --config the program shows an error and exits; for criterias.json in "
        "the startup folder it shows a warning and starts with the default criteria.",
    ])

    d.h2("13. Troubleshooting the advanced settings")
    d.table(["Message / situation", "What to do"], [
        ["Command line error, the program does not start", "Correct the parameters listed in the error message "
                                                           "(see 10.1). For their description use --help."],
        ["Warning at start", "The criterias.json or locals.json in the startup folder is invalid. The program "
                             "runs with the default settings."],
        ["English text in another language", "The text is missing from that language in locals.json - add the "
                                             "missing key."],
        ["A key such as prop.width is shown instead of a text", "The text is missing from English as well - add "
                                                                "it."],
        ["Strange characters instead of accented letters", "Save locals.json with UTF-8 encoding."],
    ], widths=[170, 300])
    d.save("TrademarkImage_User_Manual_EN")


if __name__ == "__main__":
    for doc_class in (WordDoc, MarkdownDoc):
        technical_en(doc_class)
        user_guide_hu(doc_class)
        user_guide_en(doc_class)
